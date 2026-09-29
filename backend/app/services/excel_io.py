"""Shared Excel import/export logic for the Pisell "All records" stock
export format — used by both the standalone CLI script
(scripts/import_from_excel.py) and the admin panel's one-click
export/import endpoints, so the two paths can't drift apart.

Column shape (as exported by Pisell, and mirrored on export so a
downloaded file round-trips back through import): product_id,
variant_id, title, image, vendor, category, status, stock_quantity,
price, original_price, description, variant_item1, variant_item2,
variant_item3. Each product has one parent row (variant_id blank) plus
one row per variant below it — see import_from_excel.py's module
docstring for the full parsing rationale (0-value price/stock quirks,
unpriced products, etc).
"""

import io
from collections import defaultdict

import openpyxl

from app.models.product import Product, ProductVariant
from app.services.catalog_sync import CATEGORY_MAP, ProductInput, VariantInput, to_decimal

EXTERNAL_SOURCE = "pisell"

EXPORT_COLUMNS = [
    "product_id",
    "variant_id",
    "title",
    "image",
    "vendor",
    "category",
    "status",
    "stock_quantity",
    "price",
    "original_price",
    "description",
    "variant_item1",
    "variant_item2",
    "variant_item3",
]

# Reverse of catalog_sync's Chinese-name -> CategoryId map, so a
# round-tripped export re-imports into the same category it started in
# instead of falling back to the default (map_category only recognizes
# Chinese names, not our internal ids).
CATEGORY_REVERSE_MAP: dict[str, str] = {}
for _cn_name, _cat_id in CATEGORY_MAP.items():
    CATEGORY_REVERSE_MAP.setdefault(_cat_id, _cn_name)


def load_rows(source) -> list[dict]:
    """`source` is anything openpyxl.load_workbook accepts: a file path or
    a file-like object (e.g. io.BytesIO of an uploaded file)."""
    wb = openpyxl.load_workbook(source, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    cols = [c.value for c in ws[1]]
    return [dict(zip(cols, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def first_image(raw: str | None) -> str | None:
    if not raw:
        return None
    return raw.split(";")[0].strip() or None


def variant_name(row: dict) -> str:
    """Combines the option-dimension columns into one label, e.g. a size
    and a flavor become "单个 · 柿柿如意". Falls back to the variant id if
    somehow neither dimension is set."""
    parts = []
    for item_col in ("variant_item1", "variant_item2", "variant_item3"):
        value = row.get(item_col)
        if value:
            parts.append(str(value))
    return " · ".join(parts) if parts else f"#{row.get('variant_id')}"


def build_products(rows: list[dict]) -> list[ProductInput]:
    by_product: dict[object, list[dict]] = defaultdict(list)
    for row in rows:
        pid = row.get("product_id")
        if pid is not None:
            by_product[pid].append(row)

    items: list[ProductInput] = []
    for product_id, product_rows in by_product.items():
        parent = next(
            (r for r in product_rows if not r.get("variant_id")), product_rows[0]
        )
        if not parent.get("title"):
            # Shouldn't happen, but skip rather than import a nameless product.
            continue

        price = to_decimal(parent.get("price"))
        if price is None:
            variant_prices = [
                to_decimal(r.get("price")) for r in product_rows if r is not parent
            ]
            variant_prices = [p for p in variant_prices if p]
            price = min(variant_prices) if variant_prices else None
        # Genuinely unpriced on Pisell's side (happens — a product created
        # without ever setting a price). Don't publish a $0 item; keep the
        # row (so stock/description still sync) but hide it from the store.
        is_unpriced = price is None
        if is_unpriced:
            price = to_decimal(0) or 0

        variant_rows = [r for r in product_rows if r is not parent]
        if variant_rows:
            # The parent row's own stock_quantity is unreliable — same issue
            # as price, it's sometimes 0 while the real stock sits on each
            # variant. Summing the variants is correct either way: when the
            # parent value *is* accurate, it already equals this sum.
            stock = sum((r.get("stock_quantity") or 0) for r in variant_rows)
        else:
            stock = parent.get("stock_quantity") or 0
        category = parent.get("category")

        variants = [
            VariantInput(
                external_id=str(r.get("variant_id")),
                name=variant_name(r),
                price=to_decimal(r.get("price")) or price,
                stock_quantity=int(r.get("stock_quantity") or 0),
                sort=i,
            )
            for i, r in enumerate(variant_rows)
        ]

        items.append(
            ProductInput(
                external_id=str(product_id),
                external_source=EXTERNAL_SOURCE,
                name=parent["title"],
                price=price,
                original_price=to_decimal(parent.get("original_price")),
                stock_quantity=int(stock),
                category_names=[category] if category else [],
                vendor=parent.get("vendor"),
                image_url=first_image(parent.get("image")),
                description=parent.get("description") or None,
                is_active=parent.get("status") == "published" and not is_unpriced,
                variants=variants,
            )
        )
    return items


def _product_id_value(product: Product):
    if product.external_id and product.external_id.isdigit():
        return int(product.external_id)
    return product.external_id or str(product.id)


def _variant_id_value(variant: ProductVariant):
    if variant.external_id and variant.external_id.isdigit():
        return int(variant.external_id)
    return variant.external_id or str(variant.id)


def export_workbook(products: list[Product]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(EXPORT_COLUMNS)

    for product in products:
        pid = _product_id_value(product)
        category_cn = CATEGORY_REVERSE_MAP.get(product.category, product.category)
        ws.append(
            [
                pid,
                None,
                product.name,
                product.image_url,
                product.vendor,
                category_cn,
                "published" if product.is_active else "unpublished",
                product.stock_quantity,
                float(product.price),
                float(product.old_price) if product.old_price else None,
                product.description,
                None,
                None,
                None,
            ]
        )
        for variant in product.variants:
            ws.append(
                [
                    pid,
                    _variant_id_value(variant),
                    None,
                    None,
                    None,
                    None,
                    None,
                    variant.stock_quantity,
                    float(variant.price),
                    None,
                    None,
                    variant.name,
                    None,
                    None,
                ]
            )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
